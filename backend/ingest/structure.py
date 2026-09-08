"""Stage 1 — structural analysis. Deterministic, no LLM.

Takes a merge-filled grid and works out what a human sees when they look at the
sheet: where the header row actually is, where the data starts and stops, which
columns are junk, and whether the sheet is really holding more than one table.

Every decision appends a human-readable line to the table's `notes`, because a
mis-detected header row is the single most damaging failure in the pipeline and
it has to be explainable on the review screen.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl.utils import get_column_letter

from ingest.reader import RawSheet, is_blank

# How far into a block to look for the header row.
HEADER_SCAN_ROWS = 20
# Below this, a block is not considered to start a new table.
HEADER_SCORE_FLOOR = 0.15
# Rows sampled when judging whether the data under a candidate header is
# type-consistent.
CONSISTENCY_SAMPLE = 20

SUMMARY_LABEL_RE = re.compile(
    r"^\s*(grand\s+)?(total|totals|sum|subtotal|sub-total|average|avg|count|"
    r"toplam|genel\s+toplam|ara\s+toplam)\b",
    re.IGNORECASE,
)

_NUMERIC_STR_RE = re.compile(r"^[\s$€£₺]*-?[\d.,]+\s*%?$")

# A label that introduces a block of prose rather than data. Turkish first --
# these are Turkish workbooks -- then the English equivalents.
NOTE_LABEL_RE = re.compile(
    r"^\s*(a[çc][ıi]klama(lar)?|not(lar)?|d[iı]kkat|uyar[ıi]|"
    r"note(s)?|remark(s)?|comment(s)?|caution|warning)\b\s*:?\s*$",
    re.IGNORECASE,
)
# Prose long enough that no one would call it a cell value.
PROSE_LENGTH = 40


@dataclass
class Column:
    header: str
    letter: str
    index: int          # 0-based index into the grid
    values: list[Any]


@dataclass
class DetectedSection:
    """A block of a sheet that is not table rows.

    Real workbooks end with things a table cannot hold: an AÇIKLAMALAR block
    explaining why line 5 stopped twice last night, a TOPLAM row, a report date
    floating above the header. Forcing them into columns is how you get a
    "Machine Capacity" column whose every value is the same paragraph of
    Turkish prose -- which is exactly what happened before this existed.

    They are not noise, and they are not rows. They are sections, and the
    generated app renders them under the table they came from.
    """

    kind: str              # "note" | "summary" | "meta"
    title: str
    body: str
    first_row: int         # 0-based grid rows, inclusive
    last_row: int
    source_range: str = ""


@dataclass
class DetectedTable:
    """One table. A worksheet holding two tables produces two of these."""

    name: str
    sheet_index: int
    header_row: int        # 0-based grid row
    first_data_row: int
    last_data_row: int     # inclusive
    columns: list[Column]
    header_score: float = 0.0
    notes: list[str] = field(default_factory=list)
    #: Blocks of the sheet that are not rows of this table but belong with it.
    sections: list["DetectedSection"] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return max(0, self.last_data_row - self.first_data_row + 1)

    @property
    def range_a1(self) -> str:
        if not self.columns:
            return ""
        first = self.columns[0].index
        last = self.columns[-1].index
        return (
            f"{get_column_letter(first + 1)}{self.header_row + 1}:"
            f"{get_column_letter(last + 1)}{self.last_data_row + 1}"
        )


# --------------------------------------------------------------------------
# Header detection
# --------------------------------------------------------------------------


def _looks_numeric(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        return bool(_NUMERIC_STR_RE.match(value.strip())) and any(ch.isdigit() for ch in value)
    return False


def _is_label(value: Any) -> bool:
    """A header cell is text that is not merely a number written as text."""
    return isinstance(value, str) and value.strip() != "" and not _looks_numeric(value)


def score_header_row(row: list[Any]) -> float:
    """Score a row's plausibility as a header.

    A header row is dense, made of distinct text labels, and contains no dates
    or numbers. The product of those three ratios separates it from data rows
    far more reliably than any single one of them.
    """
    non_empty = [v for v in row if not is_blank(v)]
    if len(non_empty) < 2:
        return 0.0

    non_empty_ratio = len(non_empty) / len(row)
    label_ratio = sum(1 for v in non_empty if _is_label(v)) / len(non_empty)
    normalised = [str(v).strip().lower() for v in non_empty]
    distinct_ratio = len(set(normalised)) / len(non_empty)

    return non_empty_ratio * label_ratio * distinct_ratio


def _cell_kind(value: Any) -> str:
    if is_blank(value):
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    if _looks_numeric(value):
        return "number"
    return "text"


def _consistency_below(grid: list[list[Any]], header_row: int, end_row: int) -> float:
    """How type-consistent are the rows beneath this candidate header?

    Real data columns hold one kind of value. This is the tie-breaker between
    two rows that both look header-ish (e.g. a title row and the real header).
    """
    start = header_row + 1
    stop = min(end_row + 1, start + CONSISTENCY_SAMPLE)
    if start >= stop:
        return 0.0

    width = len(grid[header_row])
    scores: list[float] = []
    for c in range(width):
        kinds = [
            _cell_kind(grid[r][c])
            for r in range(start, stop)
            if c < len(grid[r]) and not is_blank(grid[r][c])
        ]
        if not kinds:
            continue
        dominant = max(set(kinds), key=kinds.count)
        scores.append(kinds.count(dominant) / len(kinds))
    return sum(scores) / len(scores) if scores else 0.0


def find_header_row(grid: list[list[Any]], start: int, end: int) -> tuple[int, float]:
    """Return (row index, score) of the best header candidate in [start, end]."""
    best_row, best_score = start, -1.0
    limit = min(end, start + HEADER_SCAN_ROWS - 1)
    for r in range(start, limit + 1):
        base = score_header_row(grid[r])
        if base <= 0:
            continue
        # Weight mostly on the row's own shape, but let the data beneath break ties.
        score = base * (0.75 + 0.25 * _consistency_below(grid, r, end))
        if score > best_score:
            best_row, best_score = r, score
    return best_row, max(best_score, 0.0)


# --------------------------------------------------------------------------
# Block splitting (multiple tables on one sheet)
# --------------------------------------------------------------------------


def _row_is_blank(row: list[Any]) -> bool:
    return all(is_blank(v) for v in row)


def _row_blocks(grid: list[list[Any]]) -> list[tuple[int, int]]:
    """Contiguous runs of non-blank rows, as inclusive (start, end) pairs."""
    blocks: list[tuple[int, int]] = []
    start: int | None = None
    for r, row in enumerate(grid):
        if _row_is_blank(row):
            if start is not None:
                blocks.append((start, r - 1))
                start = None
        elif start is None:
            start = r
    if start is not None:
        blocks.append((start, len(grid) - 1))
    return blocks


# --------------------------------------------------------------------------
# Trailing summary rows
# --------------------------------------------------------------------------


def _is_summary_row(grid: list[list[Any]], r: int, col_span: range) -> bool:
    """A totals/subtotal row tacked on the bottom of a table.

    Two signals: an explicit label in the first populated cell, or a mostly
    empty row that still carries a number (the classic '  Total  1,234' line).
    """
    cells = [grid[r][c] for c in col_span if c < len(grid[r])]
    non_empty = [v for v in cells if not is_blank(v)]
    if not non_empty:
        return True

    first = next((v for v in cells if not is_blank(v)), None)
    if isinstance(first, str) and SUMMARY_LABEL_RE.match(first):
        return True

    blank_ratio = 1 - (len(non_empty) / len(cells))
    has_number = any(_cell_kind(v) == "number" for v in non_empty)
    return blank_ratio > 0.5 and has_number


# --------------------------------------------------------------------------
# Blocks that are not tables
# --------------------------------------------------------------------------


def _row_values(row: list[Any]) -> list[Any]:
    return [v for v in row if not is_blank(v)]


def _is_merge_filled(row: list[Any]) -> bool:
    """One value repeated across the row: a merged cell, forward-filled.

    `ingest.reader` fills every cell of a merged range with the anchor's value,
    so a paragraph merged across A:I arrives as nine identical strings. That is
    the strongest available signal that a row is one piece of prose rather than
    nine values.
    """
    values = _row_values(row)
    if len(values) < 2:
        return False
    return len({str(v).strip() for v in values}) == 1


def _is_prose_row(row: list[Any]) -> bool:
    values = _row_values(row)
    if not values:
        return False
    if _is_merge_filled(row):
        return True
    # A single long text cell in an otherwise empty row.
    return (
        len(values) == 1
        and isinstance(values[0], str)
        and len(values[0].strip()) >= PROSE_LENGTH
    )


def _has_summary_label(row: list[Any]) -> bool:
    return any(
        isinstance(v, str) and SUMMARY_LABEL_RE.match(v) for v in _row_values(row)
    )


def _is_sparse_row(row: list[Any], col_span: range) -> bool:
    """Too few populated cells to be a row of this table."""
    width = max(len(col_span), 1)
    return len(_row_values(row)) / width <= 0.5


def _is_note_label(row: list[Any]) -> bool:
    values = _row_values(row)
    return (
        len(values) == 1
        and isinstance(values[0], str)
        and bool(NOTE_LABEL_RE.match(values[0]))
    ) or (
        _is_merge_filled(row)
        and isinstance(values[0], str)
        and bool(NOTE_LABEL_RE.match(values[0]))
    )


def _block_text(grid: list[list[Any]], start: int, end: int) -> list[str]:
    """The distinct text of a block, in order.

    A block merged across columns *and* down rows arrives as the same string in
    every cell of the rectangle, so this collapses it back to the one paragraph
    the author actually wrote.
    """
    lines: list[str] = []
    for r in range(start, end + 1):
        seen_in_row: list[str] = []
        for value in _row_values(grid[r]):
            text = str(value).strip()
            if text and text not in seen_in_row:
                seen_in_row.append(text)
        for text in seen_in_row:
            if not lines or lines[-1] != text:
                lines.append(text)
    return lines


def _range(grid: list[list[Any]], start: int, end: int) -> str:
    width = max((len(grid[r]) for r in range(start, end + 1)), default=1)
    return f"A{start + 1}:{get_column_letter(max(width, 1))}{end + 1}"


def classify_block(
    grid: list[list[Any]], start: int, end: int, col_span: range
) -> DetectedSection | None:
    """Is this block prose, a totals block, or neither?

    Returns None when the block should be left to the table logic -- either as a
    table of its own or as a continuation of the one above.
    """
    rows = list(range(start, end + 1))
    prose = [r for r in rows if _is_prose_row(grid[r])]
    labels = [r for r in rows if _is_note_label(grid[r])]

    # Prose wins over everything: a paragraph merged across the sheet is not a
    # row of nine values however much it looks like one to a column counter.
    if prose or labels:
        if len(prose) + len(labels) >= max(1, len(rows) // 2):
            lines = _block_text(grid, start, end)
            title = "Notes"
            if lines and NOTE_LABEL_RE.match(lines[0]):
                title = lines[0].rstrip(" :").strip() or "Notes"
                lines = lines[1:]
            return DetectedSection(
                kind="note",
                title=title,
                body="\n\n".join(lines).strip(),
                first_row=start,
                last_row=end,
                source_range=_range(grid, start, end),
            )

    # A totals block is not always uniformly "summary rows". Merit's is three
    # rows: a PVC/TPE label row, then TOPLAM (KG), then GENEL TOPLAM (KG). What
    # identifies it is an explicit total label somewhere in the block, with
    # every row too sparse to be table data.
    labelled = any(_has_summary_label(grid[r]) for r in rows)
    sparse = all(
        _is_summary_row(grid, r, col_span) or _is_sparse_row(grid[r], col_span)
        for r in rows
    )
    if labelled and sparse:
        lines = _block_text(grid, start, end)
        return DetectedSection(
            kind="summary",
            title="Totals",
            body="\n".join(lines).strip(),
            first_row=start,
            last_row=end,
            source_range=_range(grid, start, end),
        )
    return None


# --------------------------------------------------------------------------
# Assembly
# --------------------------------------------------------------------------


def _build_columns(
    grid: list[list[Any]], header_row: int, first_data: int, last_data: int
) -> tuple[list[Column], int]:
    """Materialise columns, dropping the ones that are junk.

    A column is junk only when it has neither a header nor any data. A column
    with data but no header is kept under a generated name — throwing away real
    data because someone forgot a header would be worse than an ugly name.
    """
    width = max(
        len(grid[header_row]),
        max((len(grid[r]) for r in range(first_data, last_data + 1)), default=0),
    )

    columns: list[Column] = []
    dropped = 0
    unnamed = 0
    for c in range(width):
        header = grid[header_row][c] if c < len(grid[header_row]) else None
        values = [
            grid[r][c] if c < len(grid[r]) else None for r in range(first_data, last_data + 1)
        ]
        has_data = any(not is_blank(v) for v in values)

        if is_blank(header) and not has_data:
            dropped += 1
            continue

        if is_blank(header):
            unnamed += 1
            name = f"unnamed_{unnamed}"
        else:
            name = str(header).strip()

        columns.append(
            Column(header=name, letter=get_column_letter(c + 1), index=c, values=values)
        )
    return columns, dropped


def _analyse_block(
    sheet: RawSheet, start: int, end: int, label: str
) -> tuple[DetectedTable | None, list[str]]:
    grid = sheet.grid
    notes: list[str] = []

    header_row, score = find_header_row(grid, start, end)
    if score < HEADER_SCORE_FLOOR:
        return None, [
            f"Rows {start + 1}-{end + 1}: no row scored as a header "
            f"(best {score:.2f} < {HEADER_SCORE_FLOOR}); treated as loose rows."
        ]

    if header_row > start:
        notes.append(
            f"Skipped {header_row - start} row(s) above the header "
            f"(title/notes rows at spreadsheet row {start + 1}-{header_row})."
        )
    notes.append(f"Header detected at spreadsheet row {header_row + 1} (score {score:.2f}).")

    first_data = header_row + 1
    while first_data <= end and _row_is_blank(grid[first_data]):
        first_data += 1
    if first_data > end:
        return None, notes + [f"Header at row {header_row + 1} has no data beneath it."]

    header_cells = grid[header_row]
    col_span = range(len(header_cells))

    last_data = end
    trimmed = 0
    while last_data >= first_data and _is_summary_row(grid, last_data, col_span):
        last_data -= 1
        trimmed += 1
    if trimmed:
        notes.append(f"Trimmed {trimmed} trailing total/summary row(s).")
    if last_data < first_data:
        return None, notes + ["Every row beneath the header looked like a summary row."]

    columns, dropped = _build_columns(grid, header_row, first_data, last_data)
    if not columns:
        return None, notes + ["No usable columns after dropping empty ones."]
    if dropped:
        notes.append(f"Dropped {dropped} entirely empty column(s).")

    table = DetectedTable(
        name=label,
        sheet_index=sheet.index,
        header_row=header_row,
        first_data_row=first_data,
        last_data_row=last_data,
        columns=columns,
        header_score=score,
        notes=notes,
    )
    return table, notes


def detect_tables(sheet: RawSheet) -> list[DetectedTable]:
    """Stage 1 entry point: one worksheet in, one or more tables out."""
    if not sheet.grid:
        return []

    blocks = _row_blocks(sheet.grid)
    tables: list[DetectedTable] = []
    orphan_notes: list[str] = []

    sections: list[DetectedSection] = []

    for start, end in blocks:
        label = sheet.name if not tables else f"{sheet.name} ({len(tables) + 1})"

        # Prose and totals blocks are recognised *before* the table logic gets
        # to them. `_looks_like_continuation` asks only whether the rows put
        # something in the table's columns, which an AÇIKLAMALAR paragraph
        # merged across the sheet does trivially -- so without this check the
        # notes were absorbed as twenty data rows and every column of them read
        # as the same sentence.
        col_span = range(len(sheet.grid[start]))
        if tables:
            col_span = range(max((c.index for c in tables[-1].columns), default=0) + 1)
        section = classify_block(sheet.grid, start, end, col_span)
        if section is not None:
            sections.append(section)
            continue

        table, notes = _analyse_block(sheet, start, end, label)
        if table is None:
            # Not a table of its own. If it sits under an existing table with a
            # compatible width, it is almost certainly a continuation of it
            # after a cosmetic blank row rather than something new.
            if tables and _looks_like_continuation(sheet, tables[-1], start, end):
                previous = tables[-1]
                previous.last_data_row = end
                _refresh_columns(sheet, previous)
                previous.notes.append(
                    f"Absorbed rows {start + 1}-{end + 1} as a continuation "
                    f"(blank separator row, but no new header)."
                )
            else:
                orphan_notes.extend(notes)
            continue
        tables.append(table)

    if len(tables) > 1:
        for table in tables:
            table.notes.insert(
                0,
                f"Sheet {sheet.name!r} holds {len(tables)} tables separated by blank "
                f"rows; this is table {tables.index(table) + 1}.",
            )
    if orphan_notes and tables:
        tables[-1].notes.extend(orphan_notes)

    # Sections belong to the table they sit with. With one table on the sheet
    # that is unambiguous; with several, each section goes to the nearest table
    # above it, which is where a reader would look for it.
    for section in sections:
        owner = None
        for table in tables:
            if table.header_row < section.first_row:
                owner = table
        if owner is None and tables:
            owner = tables[0]
        if owner is not None:
            owner.sections.append(section)
            owner.notes.append(
                f"Rows {section.first_row + 1}-{section.last_row + 1} are not table "
                f"rows ({section.kind}: {section.title!r}); kept as a section."
            )

    for table in tables:
        table.notes = list(sheet.notes) + table.notes
    return tables


def _looks_like_continuation(
    sheet: RawSheet, table: DetectedTable, start: int, end: int
) -> bool:
    """Do these orphan rows populate the same columns as the table above?"""
    indexes = [col.index for col in table.columns]
    if not indexes:
        return False
    populated = 0
    for r in range(start, end + 1):
        row = sheet.grid[r]
        if any(c < len(row) and not is_blank(row[c]) for c in indexes):
            populated += 1
    return populated == (end - start + 1)


def _refresh_columns(sheet: RawSheet, table: DetectedTable) -> None:
    columns, _ = _build_columns(
        sheet.grid, table.header_row, table.first_data_row, table.last_data_row
    )
    by_index = {c.index: c for c in columns}
    for col in table.columns:
        if col.index in by_index:
            col.values = by_index[col.index].values
