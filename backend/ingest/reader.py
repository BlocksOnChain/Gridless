"""Workbook reading: file bytes -> a materialised, merge-filled cell grid.

This is the only module that knows about openpyxl. Adding .xls / CSV / Google
Sheets later means adding another function that returns RawSheet objects; the
whole pipeline downstream of here is format-agnostic. (Not implemented in v1.)
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ingest.repair import Unrepairable, strip_broken_defined_names

# Hard ceiling on rows read from any one sheet. openpyxl's reported used range
# is often inflated by stray formatting, so grids are also trimmed of trailing
# empty rows/columns after reading.
#
# This is applied as min(ws.max_row, MAX_SCAN_ROWS) -- passing it to iter_rows
# unconditionally makes openpyxl materialise that many rows even for a 60-row
# sheet, which costs seconds per sheet.
MAX_SCAN_ROWS = 200_000


class UnreadableWorkbook(Exception):
    """The uploaded bytes are not a workbook openpyxl can read."""


#: Set when a read only succeeded after `ingest.repair` cleaned the file. Read
#: by `read_workbook` so the reason is recorded on the sheet rather than lost.
REPAIR_NOTE = (
    "This workbook had {removed} broken defined name(s) (print titles or named "
    "ranges pointing at #REF!/#N/A). Excel opens such a file; openpyxl refuses "
    "it. They were ignored for this read -- the uploaded file is unchanged and "
    "no cell data is affected."
)


def _load(source, **kwargs):
    """`load_workbook`, retried once against a repaired copy.

    Real workbooks fail to open for reasons that have nothing to do with their
    contents -- see `ingest.repair`. The first attempt is the normal path and
    costs nothing extra; the retry happens only after a failure, and only for
    the one class of damage we can fix without guessing.

    Returns (workbook, repairs) where `repairs` is 0 for a clean read.
    """
    try:
        return load_workbook(source, **kwargs), 0
    except Exception as original:
        try:
            data = _source_bytes(source)
            repaired, removed = strip_broken_defined_names(data)
        except Unrepairable:
            raise original
        except Exception:
            # A repair that fails for its own reasons must not mask the real
            # error the caller needs to see.
            raise original
        try:
            return load_workbook(io.BytesIO(repaired), **kwargs), removed
        except Exception:
            raise original


def _source_bytes(source) -> bytes:
    """The bytes behind a path or a file-like, without disturbing the caller's."""
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    position = source.tell()
    try:
        source.seek(0)
        return source.read()
    finally:
        source.seek(position)


@dataclass
class RawSheet:
    """One worksheet, materialised as a rectangular row-major grid.

    Merged cells have already been forward-filled: every cell of a merged range
    carries the anchor's value, so downstream code never has to special-case
    them. `grid` is 0-indexed; `grid[r][c]` maps to spreadsheet cell
    (row r+1, column c+1).
    """

    name: str
    index: int
    grid: list[list[Any]]
    merged_ranges: int = 0
    # Columns that are empty only because the workbook stores formulas with no
    # cached result. Recorded rather than silently emitted as a null column.
    formula_only_columns: list[int] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def n_rows(self) -> int:
        return len(self.grid)

    @property
    def n_cols(self) -> int:
        return len(self.grid[0]) if self.grid else 0

    def column(self, c: int, start: int = 0, end: int | None = None) -> list[Any]:
        end = self.n_rows if end is None else end
        return [self.grid[r][c] for r in range(start, min(end, self.n_rows))]


def is_blank(value: Any) -> bool:
    """Empty for our purposes: None, or a string that is only whitespace."""
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False


def _trim(grid: list[list[Any]]) -> list[list[Any]]:
    """Drop trailing all-blank rows and columns from openpyxl's inflated range."""
    last_row = -1
    last_col = -1
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if not is_blank(value):
                last_row = r
                if c > last_col:
                    last_col = c
    if last_row < 0:
        return []
    return [row[: last_col + 1] for row in grid[: last_row + 1]]


def _apply_merges(grid: list[list[Any]], worksheet) -> int:
    """Forward-fill each merged range with its anchor value.

    openpyxl returns the value only at the top-left cell of a merged range and
    None everywhere else, which is exactly the shape that makes a human-formatted
    sheet look like it has a column of nulls. We fill the whole range instead of
    unmerging the worksheet, so the source file is never mutated.
    """
    count = 0
    n_rows = len(grid)
    for rng in worksheet.merged_cells.ranges:
        r0, c0 = rng.min_row - 1, rng.min_col - 1
        if r0 >= n_rows or c0 >= len(grid[r0]):
            continue
        anchor = grid[r0][c0]
        if anchor is None:
            continue
        for r in range(r0, min(rng.max_row, n_rows)):
            row = grid[r]
            for c in range(c0, min(rng.max_col, len(row))):
                if row[c] is None:
                    row[c] = anchor
        count += 1
    return count


def _detect_formula_only_columns(
    formula_wb, sheet_name: str, grid: list[list[Any]]
) -> list[int]:
    """Find columns that are empty *because the formulas have no cached value*.

    openpyxl's data_only=True returns None for a formula cell when the workbook
    was written programmatically and never opened by Excel. Silently emitting an
    all-null column would be a lie, so we re-open without data_only and check
    whether those cells actually hold formulas.
    """
    if not grid:
        return []
    n_cols = max(len(r) for r in grid)
    empty_cols = [
        c for c in range(n_cols) if all(is_blank(row[c]) if c < len(row) else True for row in grid)
    ]
    if not empty_cols:
        return []

    formula_cols: set[int] = set()
    ws = formula_wb[sheet_name]
    for row in ws.iter_rows(max_row=min(len(grid), 200)):
        for c in empty_cols:
            if c < len(row):
                value = row[c].value
                if isinstance(value, str) and value.startswith("="):
                    formula_cols.add(c)
    return sorted(formula_cols)


def probe_workbook(data: bytes) -> list[str]:
    """Cheap validity check at upload time: can openpyxl open this at all?

    A `.xlsx` extension is not evidence that a file is a workbook, and finding
    out during analysis means storing a file we can never use and surfacing a
    stack trace instead of a sentence.
    """
    if not data:
        raise UnreadableWorkbook("The uploaded file is empty.")
    try:
        wb, _ = _load(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise UnreadableWorkbook(
            f"This file is not a readable .xlsx workbook ({exc.__class__.__name__}). "
            f"If it was renamed from .xls or .csv, convert it in Excel first."
        ) from exc
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()
    if not names:
        raise UnreadableWorkbook("The workbook contains no sheets.")
    return names


def read_workbook(path: str | Path) -> list[RawSheet]:
    """Read every worksheet into a merge-filled grid of *values*.

    Formula cells yield their cached value, never the formula string.
    """
    path = Path(path)
    wb, repairs = _load(path, data_only=True)
    # Opened once and shared: the formula view is only consulted for sheets that
    # actually have empty columns, but re-opening it per sheet is expensive.
    formula_wb = None
    sheets: list[RawSheet] = []
    try:
        for index, name in enumerate(wb.sheetnames):
            ws = wb[name]
            max_row = min(ws.max_row or 0, MAX_SCAN_ROWS)
            grid = (
                [list(row) for row in ws.iter_rows(max_row=max_row, values_only=True)]
                if max_row
                else []
            )
            merged = _apply_merges(grid, ws)
            grid = _trim(grid)

            sheet = RawSheet(name=name, index=index, grid=grid, merged_ranges=merged)
            if repairs:
                sheet.notes.append(REPAIR_NOTE.format(removed=repairs))
            if merged:
                sheet.notes.append(
                    f"Forward-filled {merged} merged cell range(s) so grouped labels "
                    f"appear on every row they cover."
                )
            if not grid:
                sheet.notes.append("Sheet is empty; skipped.")
                sheets.append(sheet)
                continue

            if formula_wb is None:
                formula_wb, _ = _load(path, data_only=False, read_only=True)
            formula_cols = _detect_formula_only_columns(formula_wb, name, grid)
            if formula_cols:
                letters = ", ".join(get_column_letter(c + 1) for c in formula_cols)
                sheet.formula_only_columns = formula_cols
                sheet.notes.append(
                    f"Column(s) {letters} contain formulas with no cached value "
                    f"(the file was never opened by Excel), so they read as empty. "
                    f"Open and re-save the workbook in Excel to recover them."
                )
            sheets.append(sheet)
    finally:
        wb.close()
        if formula_wb is not None:
            formula_wb.close()
    return sheets
